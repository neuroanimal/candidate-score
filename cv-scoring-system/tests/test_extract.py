"""PDF extraction: InDesign-like two-column layout, links, OCR fallback."""

import pytest

from cvscore.extract import extract, fix_ocr_line, ocr_available
from cvscore.sections import split_sections, timeline


def test_two_column_reading_order(sample_pdfs, today):
    """Two column reading order."""
    doc = extract(sample_pdfs / "cv_2col.pdf", ocr="never")
    text = doc.text
    assert doc.max_columns == 2
    # the sidebar is read as one block and the main column is not interleaved with it
    assert text.index("UMIEJĘTNOŚCI") < text.index("Prawo jazdy") < text.index("PROFIL")
    assert text.index("PROFIL") < text.index("WYKSZTAŁCENIE")
    # wrapped bullet lines are re-joined
    assert "w Photoshopie, Illustratorze i InDesignie" in text
    assert "Adobe Photoshop, Illustrator, InDesign" in text
    # metadata and hidden links
    assert doc.metadata.get("indesign")
    assert any("behance" in u for u in doc.all_links())
    jobs = [e for e in timeline(split_sections(doc), today) if e.section == "experience"]
    assert len(jobs) == 3 and jobs[0].ongoing


def test_page_has_photo_box(sample_pdfs):
    """Page has photo box."""
    doc = extract(sample_pdfs / "cv_2col.pdf", ocr="never")
    assert doc.pages[0].image_boxes


@pytest.mark.skipif(not ocr_available(), reason="Tesseract not installed")
def test_outlines_need_ocr(sample_pdfs, today):
    """Outlines need ocr."""
    doc = extract(sample_pdfs / "cv_outlines.pdf")
    assert doc.ocr_pages == 1 and doc.pages[0].ocr_confidence > 80
    assert "Prowadziłam" in doc.text and "Akademia Sztuk Pięknych" in doc.text
    jobs = [e for e in timeline(split_sections(doc), today) if e.section == "experience"]
    assert len(jobs) == 3


def test_ocr_never_on_outlines_gives_warning_free_empty_text(sample_pdfs):
    """Ocr never on outlines gives warning free empty text."""
    doc = extract(sample_pdfs / "cv_outlines.pdf", ocr="never")
    assert doc.pages[0].chars == 0


def test_fix_ocr_line():
    """Fix ocr line."""
    assert fix_ocr_line("» Prowadzitam zajęcia") == "• Prowadziłam zajęcia"
    assert fix_ocr_line("Witam serdecznie") == "Witam serdecznie"


def test_docx_and_txt(tmp_path):
    """Docx and txt."""
    import docx  # pylint: disable=import-outside-toplevel

    d = docx.Document()
    d.add_heading("Doświadczenie zawodowe", level=1)
    d.add_paragraph("03.2021 – obecnie Instruktorka, SDS")
    d.save(tmp_path / "cv.docx")
    doc = extract(tmp_path / "cv.docx")
    assert doc.kind == "docx" and "Instruktorka" in doc.text
    (tmp_path / "cv.txt").write_bytes("Doświadczenie\nŁódź".encode("cp1250"))
    assert "Łódź" in extract(tmp_path / "cv.txt").text
    with pytest.raises(ValueError):
        extract(tmp_path / "cv.odt")


def test_narrow_date_column_stays_with_entries(tmp_path, today):
    """A left column of dates is not split off as a separate column; 3 skill columns read top-down."""
    import make_samples as ms  # pylint: disable=import-outside-toplevel

    w = ms.Writer()
    w.text(50, 60, "Anna Przykładowa", 22, True)
    y = 120
    w.text(50, y, "DOŚWIADCZENIE ZAWODOWE", 12, True)
    y += 24
    for dates, title, bullet in (("2021 – obecnie", "Graficzka, Studio Przykład", "Zaprojektowałam 30 identyfikacji"),
                                 ("2018 – 2021", "Ilustratorka, Wydawnictwo", "Zilustrowałam 12 książek dla dzieci")):
        w.text(50, y, dates, 9)
        w.text(150, y, title, 10, True)
        y = w.wrap(150, y + 14, bullet, 380, 9, bullet=True) + 12
    w.text(50, y + 10, "UMIEJĘTNOŚCI", 12, True)
    for i, col in enumerate((["Photoshop", "Illustrator"], ["Linoryt", "Akwarela"], ["Instagram", "Canva"])):
        for j, skill in enumerate(col):
            w.text(50 + 170 * i, y + 30 + 13 * j, skill, 9)
    w.save(tmp_path / "datecol.pdf")
    doc = extract(tmp_path / "datecol.pdf")
    assert "2021 – obecnie Graficzka, Studio Przykład" in doc.text
    assert doc.text.index("Photoshop") < doc.text.index("Illustrator") < doc.text.index("Linoryt")
    jobs = timeline(split_sections(doc), today)
    assert [j.title for j in jobs] == ["Graficzka, Studio Przykład", "Ilustratorka, Wydawnictwo"]
    assert jobs[1].lines == ["• Zilustrowałam 12 książek dla dzieci"]


def test_rotated_text_does_not_break_layout(tmp_path):
    """A vertical label along the edge is kept but read after the page body."""
    import make_samples as ms  # pylint: disable=import-outside-toplevel

    w = ms.Writer()
    w.page.insert_text((20, 700), "PORTFOLIO 2026", fontname="dj", fontsize=14, rotate=90)
    w.text(60, 80, "DOŚWIADCZENIE ZAWODOWE", 12, True)
    w.text(60, 100, "2021 – obecnie Graficzka, Studio Przykład", 10)
    w.text(60, 114, "• Zaprojektowałam 30 identyfikacji wizualnych", 9)
    w.save(tmp_path / "rot.pdf")
    doc = extract(tmp_path / "rot.pdf", ocr="never")
    assert doc.text.index("DOŚWIADCZENIE") < doc.text.index("PORTFOLIO 2026")
    assert "2021 – obecnie Graficzka" in doc.text


def test_extraction_cache(sample_pdfs, tmp_path, monkeypatch):
    """A second extraction of the same content comes from the cache, under the new path."""
    import shutil  # pylint: disable=import-outside-toplevel
    from cvscore import extract as ex  # pylint: disable=import-outside-toplevel

    monkeypatch.setenv("CVSCORE_CACHE", "1")
    monkeypatch.setenv("CVSCORE_CACHE_DIR", str(tmp_path / "cache"))
    first = ex.extract(sample_pdfs / "cv_2col.pdf")
    copy = tmp_path / "kopia.pdf"
    shutil.copy(sample_pdfs / "cv_2col.pdf", copy)
    monkeypatch.setattr(ex, "extract_pdf", lambda *a, **k: (_ for _ in ()).throw(AssertionError("not cached")))
    second = ex.extract(copy)
    assert second.text == first.text and second.path == str(copy)
